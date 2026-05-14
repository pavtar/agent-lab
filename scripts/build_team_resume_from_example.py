#!/usr/bin/env python3
"""
Собирает презентацию резюме: клонирует слайд из «Пример Резюме.pptx»,
заполняет поля команды «Путь инженера». Без фото.
"""
from __future__ import annotations

import copy
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER


EXAMPLE = Path(
    "/Users/paveltarelkin/Desktop/Vibe/sales ai/Путь инженера/Пример Резюме.pptx"
)
OUTPUT = Path(
    "/Users/paveltarelkin/Desktop/Vibe/sales ai/Путь инженера/Резюме_команды_1.pptx"
)

APPENDIX_TITLE = "Приложение. Резюме участников команды"

# Вторая строка прямоугольника под заголовком: роль на проекте (общая для команды)
ROLE_ON_PROJECT = (
    "Участник проектной команды направления «Путь инженера», ООО «Аплэб»."
)


def clone_slide_clean(pres: Presentation, slide_index: int):
    slide = pres.slides[slide_index]
    dup = pres.slides.add_slide(slide.slide_layout)
    sp_tree = dup.shapes._spTree
    for el in list(sp_tree):
        tag = el.tag.split("}")[-1]
        if tag not in ("nvGrpSpPr", "grpSpPr"):
            sp_tree.remove(el)
    for shape in slide.shapes:
        sp_tree.append(copy.deepcopy(shape._element))
    return dup


def remove_picture_placeholder(slide) -> None:
    for shape in list(slide.shapes):
        if not shape.is_placeholder:
            continue
        try:
            if shape.placeholder_format.type == PP_PLACEHOLDER.PICTURE:
                shape.element.getparent().remove(shape.element)
        except Exception:
            pass


def set_cell_bullets(cell, lines: list[str]) -> None:
    """Текст ячейки таблицы: маркеры «•» построчно."""
    text = "\n".join(f"• {line}" if not line.startswith("•") else line for line in lines)
    cell.text = text


def fill_slide(slide, person: dict) -> None:
    slide.shapes[0].text_frame.text = APPENDIX_TITLE

    position = person["position"]
    slide.shapes[1].text_frame.text = f"{position}\n{ROLE_ON_PROJECT}"

    slide.shapes[2].text_frame.text = person["name_if"]  # Имя Фамилия

    remove_picture_placeholder(slide)

    tbl_shape = None
    for sh in slide.shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.TABLE:
            tbl_shape = sh
            break
    if tbl_shape is None:
        raise RuntimeError("Таблица не найдена")

    tbl = tbl_shape.table
    tbl.cell(0, 0).text = "Образование, квалификация"
    set_cell_bullets(tbl.cell(0, 1), person["edu_bullets"])
    tbl.cell(1, 0).text = "Проектный опыт"
    set_cell_bullets(tbl.cell(1, 1), person["proj_bullets"])


TEAM: list[dict] = [
    {
        "name_if": "Андрей Головинов",
        "position": "Руководитель отдела бэкенд-разработки",
        "edu_bullets": [
            "Квалификация Teamlead: Node.js и экосистема, высокопроизводительные приложения.",
            "Docker, Kubernetes, микросервисы, отказоустойчивость.",
            "PostgreSQL, Redis, NoSQL.",
            "CI/CD; высоконагруженные сервисы и интеграции с AWS.",
            "Команды, Agile, менторинг; международные продукты, английский с клиентами.",
        ],
        "proj_bullets": [
            "СЕВЭНКО (СТС), Алюминиевая Ассоциация, Mitsuparts, «Стальные решения» (Северсталь).",
            "Часть проектов под NDA — по запросу демонстрация на стенде.",
            "ООО «Аплэб», руководитель отдела бэкенд-разработки (с 19.08.2025).",
        ],
    },
    {
        "name_if": "Андрей Горб",
        "position": "Руководитель отдела аналитики",
        "edu_bullets": [
            "Исследования сайтов металлургических и смежных компаний, адаптация практик.",
            "Структура сайта под бизнес и сегменты пользователей.",
            "Участие на всех этапах разработки, контроль соответствия целям.",
            "Руководство проектом: руководитель, системный и бизнес-аналитик.",
        ],
        "proj_bullets": [
            "Портал «Богослов.RU», СИБУР ПолиЛаб, Алюминиевая ассоциация компаний.",
            "ООО «Аплэб», руководитель отдела аналитики (с 16.12.2024).",
        ],
    },
    {
        "name_if": "Анна Долганова",
        "position": "Дизайнер",
        "edu_bullets": [
            "Бенчмаркинг и лучшие практики.",
            "Проектирование структуры, интерфейсов и пользовательских потоков.",
            "Айдентика, брендинг, концепты, гайдлайны.",
            "Технический дизайн; презентации и стратегические обсуждения.",
        ],
        "proj_bullets": [
            "Memini, реабилитация после инсульта, Сиалорея, карьер Логики молока.",
            "Северсталь («Стальные решения», «Вместе»), 3D-конфигуратор, GNM, Reksoft, СТС, СЕВЭНКО, аэропорт Новокузнецк, Сибур Полилаб, фасадные решения и др.",
            "ООО «Аплэб», дизайнер (с 06.02.2024).",
        ],
    },
    {
        "name_if": "Дмитрий Мекеров",
        "position": "Frontend-разработчик",
        "edu_bullets": [
            "Методики и подходы к разработке интерфейсов.",
            "Архитектура проекта и компонентов.",
            "Продуктовое тестирование web: A/B, Canary и др.",
            "Метрики: Яндекс.Метрика, Google Analytics.",
            "Core Web Vitals, оптимизация производительности, кэш, заголовки запросов.",
        ],
        "proj_bullets": [
            "МТТ, фонд «Черкизово», «Атриум», Алюминиевая Ассоциация, Folga, аэропорт Геленджик.",
            "iiii-tech, Мосполитех, «Стальное дерево», сайты Корпорации СТС.",
            "ООО «Аплэб», frontend-разработчик (с 06.04.2018).",
        ],
    },
    {
        "name_if": "Александр Мечинзов",
        "position": "Бэкенд-разработчик",
        "edu_bullets": [
            "Высшее: УАТУ, «Бизнес-информатика, прикладная информатика».",
            "Сертификаты 1С-Битрикс: администратор, разработчик, композит, интеграция дизайна, расширение возможностей.",
            "PHP, MySQL, HTML/CSS/JS; БД; Linux, SSH; Docker/Kubernetes; Git, CI/CD.",
            "1С-Битрикс, команда, код-ревью; BitrixVM, Composer, NPM, Laravel; БУС.",
        ],
        "proj_bullets": [
            "ПГК, МТТ, Kept, Саянская фольга, Элара, Окнатрейд, Вольта, карьер Полюса, iiii-tech, Флит Финанс.",
            "«Стальное дерево», Aluminas, Атриум, аэропорт Самарканд, Меллинг Войтишкин и Партнёры, СЕВЭНКО, карьер Логики молока.",
            "ООО «Аплэб», бэкенд-разработчик (с 05.07.2022).",
        ],
    },
    {
        "name_if": "Иван Петров",
        "position": "Ведущий контент-менеджер",
        "edu_bullets": [
            "Курсы: контент-менеджер; администратор 1С-Битрикс (базовый, модули); «Многосайтовость».",
            "Контент, грамотная речь, HTML/CSS, CMS, графические редакторы.",
            "1С-Битрикс, Tilda, Figma; контент-планы; товары; верстка статей; актуальность контента.",
        ],
        "proj_bullets": [
            "Атол, Атол-Онлайн, аэропорты, Вольта, Иннопром, Kept, НЛМК, Ла Моена, ММК, Мособлкино.",
            "Кейс Газ групп (Uplab), Северсталь «Вместе», Элара, EN+, МТТ, Осетр, Reef Life и др.",
            "ООО «Аплэб», ведущий контент-менеджер (с 16.07.2018).",
        ],
    },
    {
        "name_if": "Анна Питернова",
        "position": "Младший тестировщик (QA)",
        "edu_bullets": [
            "Высшее: МГТУ; курс Geek Brains «Быстрый старт. Тестировщик».",
            "Мануальное QA: требования, функционал, usability, UI, совместимость, мобильное, регресс.",
            "Админки Bitrix, WordPress, Symfony; баг-репорты, чек-листы; Jira.",
            "Agile, Scrum, Kanban.",
        ],
        "proj_bullets": [
            "Аэропорты Самарканд и Южно-Сахалинск, фасадные решения, Kept, СТС, Ла Моена, Полюс.",
            "Северсталь («Стальные решения», «Вместе»), Черкизово, ПГК.",
            "ООО «Аплэб», QA.",
        ],
    },
    {
        "name_if": "Елена Пономарёва",
        "position": "Арт-директор",
        "edu_bullets": [
            "Концепты, графика, коммуникации; ТЗ, планирование, сметы, тайминги.",
            "UI/UX, айдентика, брендинг, гайдлайны; крупные DEV-проекты.",
            "Исследования, тренды, контроль исполнителей; бенчмаркинг, техдизайн.",
        ],
        "proj_bullets": [
            "Логика молока, Черкизово, аэропорт Самарканд, GNM, серия аэропортов, Саянская фольга, СТС, Атриум.",
            "Меллинг Войтишкин и Партнёры, НЛМК «Дружный двор», Атол-Онлайн, кейс Чкалов, X5, Экомилк.",
            "ООО «Аплэб», арт-директор (с 09.2024); ранее ведущий дизайнер и дизайнер.",
        ],
    },
    {
        "name_if": "Михаил Мамедов",
        "position": "Руководитель проектной группы",
        "edu_bullets": [
            "Технический бэкграунд, высшее образование.",
            "Agile, Scrum, Kanban; PMBoK; полный цикл проекта.",
            "Confluence, Jira, диаграммы Ганта; корпоративные сайты, порталы, e-commerce.",
        ],
        "proj_bullets": [
            "rehabafterstroke.ru, memini.ru, детский проект Сиалореи, dataplatform.ru, datasapience.ru, klimona.life.",
            "Коммуникационный дизайн «Мегамаркет»; расширенный список по запросу.",
            "ООО «Аплэб», руководитель проектной группы (по портфолио с 13.04.2026).",
        ],
    },
    {
        "name_if": "Кристина Ткаченко",
        "position": "Руководитель проекта",
        "edu_bullets": [
            "Agile, Scrum, Kanban; PMBoK; веб-приложения.",
            "Confluence, Jira, диаграммы Ганта, тепловые карты.",
            "Вывод в продакшн; координация команд.",
        ],
        "proj_bullets": [
            "Черкизово, фонд Черкизово, Меллинг Войтишкин и Партнёры, «Стальные решения» Северстали.",
            "Дизайн для «Лаборатории Касперского» (NDA); расширенный список по запросу.",
            "ООО «Аплэб», руководитель проектов (с 01.08.2024).",
        ],
    },
    {
        "name_if": "Эмиль Тухватуллин",
        "position": "Ведущий Frontend-разработчик",
        "edu_bullets": [
            "Курсы: JavaScript, Node.js, Nuxt (Vue); CMS: 1C-Bitrix, WordPress, Тильда.",
            "HTML5, сложные UI, производительность, WebSocket, GraphQL/REST.",
            "Алгоритмы, паттерны; Cypress/Jest; SQL/NoSQL.",
        ],
        "proj_bullets": [
            "Aluminas, Атол, Атол-Онлайн, аэропорт Геленджик, Вольта, Kept, НЛМК «Стальное дерево», Мерц, Прайдекс.",
            "Полюс, Саянская фольга, Северсталь, Элара; остальное NDA — стенд по запросу.",
            "ООО «Аплэб», frontend (с 05.08.2019).",
        ],
    },
]


def main() -> int:
    if not EXAMPLE.is_file():
        print(f"Нет файла примера: {EXAMPLE}", file=sys.stderr)
        return 1

    prs = Presentation(str(EXAMPLE))
    slides = [prs.slides[0]]
    for _ in range(1, len(TEAM)):
        slides.append(clone_slide_clean(prs, 0))

    for slide, person in zip(slides, TEAM):
        fill_slide(slide, person)

    prs.save(str(OUTPUT))
    print(f"Saved {OUTPUT} ({len(prs.slides)} slides)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
