"""
Presentation engine (v3 — simplified).

Orchestrates rendering of SlideSpecs into PPTX.
No more ScaleContext with relative coordinates — helpers use fixed Inches.
"""

from __future__ import annotations

import os
from pptx import Presentation
from pptx.oxml.ns import qn

from . import styles as S
from .parser import SlideSpec
from . import slide_types
from . import helpers


RENDERERS = {
    "cover": slide_types.render_cover,
    "section": slide_types.render_section,
    "stats": slide_types.render_stats,
    "text": slide_types.render_text,
    "columns": slide_types.render_columns,
    "table": slide_types.render_table,
    "steps": slide_types.render_steps,
    "case": slide_types.render_case,
    "image": slide_types.render_image,
    "grid": slide_types.render_grid,
    "quote": slide_types.render_quote,
    "highlight": slide_types.render_highlight,
    "timeline": slide_types.render_timeline,
    "contacts": slide_types.render_contacts,
    "comparison": slide_types.render_table,
    "toc": slide_types.render_toc,
    "team": slide_types.render_team,
}


def generate(
    template_path: str,
    slide_specs: list[SlideSpec],
    output_path: str,
    breadcrumb: str = "",
) -> str:
    prs = Presentation(template_path)
    _remove_template_slides(prs)
    layout_map = _build_layout_map(prs)

    for i, spec in enumerate(slide_specs):
        layout = _pick_layout(spec, layout_map)
        slide = prs.slides.add_slide(layout)
        _clear_placeholders(slide)

        renderer = RENDERERS.get(spec.slide_type, slide_types.render_text)
        renderer(slide, spec)

        if spec.slide_type != "cover":
            helpers.LOGO_BOTTOM_RIGHT(slide)

        if breadcrumb and spec.slide_type not in ("cover", "section", "contacts"):
            helpers.add_breadcrumb(slide, breadcrumb)

        if spec.slide_type not in ("cover", "contacts"):
            helpers.add_slide_number(slide, i + 1)

        if spec.note:
            slide.notes_slide.notes_text_frame.text = spec.note

    prs.save(output_path)
    return os.path.abspath(output_path)


def _build_layout_map(prs: Presentation) -> dict:
    return {layout.name: layout for layout in prs.slide_layouts}


def _pick_layout(spec: SlideSpec, layouts: dict):
    preferred = {
        "cover":      [S.LAYOUT_MINT, S.LAYOUT_DEFAULT, S.LAYOUT_BLANK],
        "section":    [S.LAYOUT_MINT, S.LAYOUT_DEFAULT],
        "contacts":   [S.LAYOUT_MINT, S.LAYOUT_GRAY, S.LAYOUT_BLANK],
    }

    for name in preferred.get(spec.slide_type, [S.LAYOUT_BLANK, S.LAYOUT_DEFAULT]):
        if name in layouts:
            return layouts[name]

    return next(iter(layouts.values()))


def _remove_template_slides(prs: Presentation):
    """Delete every existing slide, keeping only layouts/masters."""
    lst = prs.slides._sldIdLst
    for sld_id in list(lst):
        lst.remove(sld_id)
    to_drop = [
        rel for rel in prs.part.rels.values()
        if "slide" in rel.reltype
        and "slideLayout" not in rel.reltype
        and "slideMaster" not in rel.reltype
    ]
    for rel in to_drop:
        prs.part.drop_rel(rel.rId)


def _clear_placeholders(slide):
    """Remove inherited placeholder shapes so we render from scratch."""
    for ph in list(slide.placeholders):
        ph._element.getparent().remove(ph._element)
